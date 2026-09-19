import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def create_manuscript():
    doc = docx.Document()

    # Margins
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Base Normal Style
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(10.5)
    font.color.rgb = RGBColor(0x22, 0x22, 0x22)

    # Title
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run('PRIE: An Explainable Placement Readiness Intelligence Engine with Calibrated Predictive Modeling and Constrained Prescriptive Recourse')
    title_run.font.size = Pt(18)
    title_run.font.bold = True
    title_run.font.name = 'Times New Roman'
    title_p.paragraph_format.space_after = Pt(12)

    # Authors
    author_p = doc.add_paragraph()
    author_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    a_run = author_p.add_run('Anonymous Authors\nDepartment of Computer Science and Engineering\nAffiliated Engineering Institution, City, Country\nemail@institution.edu\n')
    a_run.font.size = Pt(10)
    a_run.font.italic = True
    author_p.paragraph_format.space_after = Pt(18)

    with open('10_Publication/01_Conference_Paper/paper_manuscript.md', encoding='utf-8') as f:
        text = f.read()

    lines = text.split('\n')
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue

        if line.startswith('# PRIE:') or line.startswith('**Anonymous Authors**') or 'email@institution.edu' in line or line == '---':
            i += 1
            continue

        if line.startswith('## Abstract'):
            p = doc.add_heading('Abstract', level=1)
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            i += 1
            abs_text = ''
            while i < len(lines) and not lines[i].startswith('**Keywords**') and not lines[i].startswith('## '):
                if lines[i].strip():
                    abs_text += lines[i].strip() + ' '
                i += 1
            p_abs = doc.add_paragraph()
            r_abs = p_abs.add_run(abs_text.strip())
            r_abs.font.size = Pt(9.5)
            r_abs.font.bold = True
            p_abs.paragraph_format.space_after = Pt(8)

            if i < len(lines) and lines[i].startswith('**Keywords**'):
                p_kw = doc.add_paragraph()
                r_kw = p_kw.add_run(lines[i].strip())
                r_kw.font.size = Pt(9.5)
                r_kw.font.italic = True
                p_kw.paragraph_format.space_after = Pt(14)
                i += 1
            continue

        if line.startswith('## '):
            heading_text = line.replace('## ', '').strip()
            p = doc.add_heading(heading_text, level=1)
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            i += 1
            continue

        if line.startswith('### '):
            heading_text = line.replace('### ', '').strip()
            p = doc.add_heading(heading_text, level=2)
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            i += 1
            continue

        if line.startswith('|') and '|' in line[1:]:
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                table_lines.append(lines[i].strip())
                i += 1

            headers = [c.strip() for c in table_lines[0].split('|')[1:-1]]
            data_rows = []
            for r in table_lines[2:]:
                cols = [c.strip() for c in r.split('|')[1:-1]]
                if len(cols) == len(headers):
                    data_rows.append(cols)

            table = doc.add_table(rows=len(data_rows) + 1, cols=len(headers))
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            for col_idx, h in enumerate(headers):
                cell = table.cell(0, col_idx)
                cell.text = h.replace('**', '')
                if cell.paragraphs[0].runs:
                    cell.paragraphs[0].runs[0].font.bold = True
                    cell.paragraphs[0].runs[0].font.size = Pt(9)
                shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="EAEAEA"/>')
                cell._tc.get_or_add_tcPr().append(shading_elm)

            for row_idx, r_data in enumerate(data_rows):
                for col_idx, val in enumerate(r_data):
                    cell = table.cell(row_idx + 1, col_idx)
                    cell.text = val.replace('**', '').replace('\\%', '%')
                    if cell.paragraphs[0].runs:
                        cell.paragraphs[0].runs[0].font.size = Pt(8.5)

            p_sp = doc.add_paragraph()
            p_sp.paragraph_format.space_after = Pt(8)
            continue

        if line.startswith('!['):
            caption = line[line.find('[')+1:line.find(']')]
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(f'[{caption}]')
            r.font.bold = True
            r.font.italic = True
            r.font.size = Pt(9)
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(8)
            i += 1
            continue

        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        if line.startswith('* ') or line.startswith('- '):
            p.style = 'List Bullet'
            p.add_run(line[2:].strip())
        elif line[0].isdigit() and (line[1:3] == '. ' or line[2:4] == '. '):
            p.style = 'List Number'
            dot_idx = line.find('. ')
            p.add_run(line[dot_idx+2:].strip())
        else:
            p.add_run(line)
        i += 1

    doc.save('10_Publication/01_Conference_Paper/paper.docx')
    doc.save('10_Publication/Conference_Paper/paper.docx')
    print('SUCCESS: Created paper.docx in 01_Conference_Paper and Conference_Paper')

if __name__ == '__main__':
    create_manuscript()
