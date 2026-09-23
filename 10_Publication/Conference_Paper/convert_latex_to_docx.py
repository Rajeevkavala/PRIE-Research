import os
import re
import shutil
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def clean_latex_text(text, bib_map):
    if not text:
        return ""
    
    # Replace citations \cite{b1,b2} with [1, 2]
    def replace_cite(m):
        keys = [k.strip() for k in m.group(1).split(',')]
        nums = [str(bib_map.get(k, k)) for k in keys]
        return "[" + ", ".join(nums) + "]"
    text = re.sub(r'\\cite\{([^}]+)\}', replace_cite, text)
    
    # Mathematical symbols and greek letters
    text = text.replace(r'\le', '≤')
    text = text.replace(r'\ge', '≥')
    text = text.replace(r'\pm', '±')
    text = text.replace(r'\Delta', 'Δ')
    text = text.replace(r'\tau', 'τ')
    text = text.replace(r'\phi_i', 'φ_i')
    text = text.replace(r'\phi', 'φ')
    text = text.replace(r'\in', '∈')
    text = text.replace(r'\times', '×')
    text = text.replace(r'\rightarrow', '→')
    text = text.replace(r'\approx', '≈')
    text = text.replace(r'\neq', '≠')
    text = text.replace(r'\sim', '~')
    text = text.replace(r'\cdot', '·')
    
    # Common LaTeX formatting
    text = text.replace(r'\%', '%')
    text = text.replace(r'\&', '&')
    text = text.replace(r'\$', '$')
    text = text.replace(r'\_', '_')
    text = text.replace(r'\,', ' ')
    text = text.replace(r'\ ', ' ')
    text = text.replace('~', ' ')
    text = text.replace('---', '—')
    text = text.replace('--', '–')
    text = text.replace("``", '“').replace("''", '”')
    text = text.replace("`", "‘").replace("'", "’")
    
    # Strip \texttt, \textit, \textbf, \text, \mathrm
    text = re.sub(r'\\texttt\{([^}]+)\}', r'\1', text)
    text = re.sub(r'\\textit\{([^}]+)\}', r'\1', text)
    text = re.sub(r'\\textbf\{([^}]+)\}', r'\1', text)
    text = re.sub(r'\\text\{([^}]+)\}', r'\1', text)
    text = re.sub(r'\\mathrm\{([^}]+)\}', r'\1', text)
    text = re.sub(r'\\mathcal\{([^}]+)\}', r'\1', text)
    text = re.sub(r'\\mathbf\{([^}]+)\}', r'\1', text)
    
    # Specific math cleanup
    text = text.replace(r'x_{\text{spv}}', 'x_spv')
    text = text.replace(r'x_{\mathrm{spv}}', 'x_spv')
    text = text.replace(r'x_{spv}', 'x_spv')
    text = text.replace(r'\hat{p}', 'p̂')
    text = text.replace(r'z(x)', 'z(x)')
    text = text.replace(r'\{42, 123, 456, 789, 2026\}', '{42, 123, 456, 789, 2026}')
    text = text.replace(r'\{0, 1\}^{22}', '{0, 1}^22')
    text = text.replace(r'[0.0, 1.0]^{22}', '[0.0, 1.0]^22')
    text = text.replace(r'O(TLD^2)', 'O(TLD²)')
    text = text.replace(r'w_1 = 0.53', 'w₁ = 0.53')
    text = text.replace(r'w_0 = 1.0', 'w₀ = 1.0')
    text = text.replace(r'\lambda', 'λ')
    
    # Strip lingering math dollar signs
    text = text.replace('$', '')
    
    # Clean up double spaces
    text = re.sub(r'\s+', ' ', text).strip()
    return text

def set_table_borders_ieee(table):
    tblPr = table._tbl.tblPr
    borders_xml = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="single" w:sz="12" w:space="0" w:color="000000"/>\n'
        f'  <w:left w:val="none"/>\n'
        f'  <w:bottom w:val="single" w:sz="12" w:space="0" w:color="000000"/>\n'
        f'  <w:right w:val="none"/>\n'
        f'  <w:insideH w:val="none"/>\n'
        f'  <w:insideV w:val="none"/>\n'
        f'</w:tblBorders>'
    )
    tblPr.append(borders_xml)

def set_header_bottom_border(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    b_xml = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:bottom w:val="single" w:sz="6" w:space="0" w:color="000000"/></w:tcBorders>')
    tcPr.append(b_xml)

def convert_tex_to_docx(tex_path, docx_out, txt_out):
    print(f"Reading {tex_path}...")
    with open(tex_path, 'r', encoding='utf-8') as f:
        tex = f.read()

    # Build bibliography mapping
    bib_map = {}
    bib_items = []
    bib_matches = list(re.finditer(r'\\bibitem\{([^}]+)\}\s*([^\n\\]+(?:\\textit\{[^}]+\}[^\n\\]*|\\text\{[^}]+\}[^\n\\]*|\\textbf\{[^}]+\}[^\n\\]*|(?!\s*\\bibitem|\s*\\end\{thebibliography\})[^\\]*)*)', tex))
    for idx, bm in enumerate(re.finditer(r'\\bibitem\{([^}]+)\}(.*?)(?=\\bibitem|\s*\\end\{thebibliography\})', tex, re.DOTALL), 1):
        key = bm.group(1).strip()
        body = bm.group(2).strip()
        bib_map[key] = idx
        cleaned_body = clean_latex_text(body, bib_map)
        bib_items.append((idx, key, cleaned_body))

    doc = docx.Document()

    # Section 0: Title, Authors, Abstract, Keywords (Single Column Header)
    s0 = doc.sections[0]
    s0.top_margin = Inches(0.75)
    s0.bottom_margin = Inches(0.75)
    s0.left_margin = Inches(0.75)
    s0.right_margin = Inches(0.75)
    s0.page_width = Inches(8.5)
    s0.page_height = Inches(11.0)

    # Title
    m_title = re.search(r'\\title\{([^}]+)\}', tex)
    title_text = m_title.group(1) if m_title else "PRIE: Placement Readiness Intelligence Engine"
    title_text = clean_latex_text(title_text, bib_map)

    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run(title_text)
    title_run.font.size = Pt(18)
    title_run.font.bold = True
    title_run.font.name = 'Times New Roman'
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(12)

    # Authors Table
    authors_grid = [
        [
            ("Sivasubramanian R", "Department of AIML\nMalla Reddy University\nDulapally, Hyderabad, India\nSivasubramanian243@gmail.com"),
            ("Kavala Rajeev", "Department of AIML\nMalla Reddy University\nDulapally, Hyderabad, India\nKavalarajeev@gmail.com")
        ],
        [
            ("Kundala Dhana Naga Shankar", "Department of AIML\nMalla Reddy University\nDulapally, Hyderabad, India\nKundaladhana2004@gmail.com"),
            ("Kouru Rudra Teja", "Department of AIML\nMalla Reddy University\nDulapally, Hyderabad, India\nRudrateja08@gmail.com")
        ]
    ]

    t_auth = doc.add_table(rows=2, cols=2)
    t_auth.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r_idx in range(2):
        row = t_auth.rows[r_idx]
        for c_idx in range(2):
            name, affil = authors_grid[r_idx][c_idx]
            cell = row.cells[c_idx]
            cell.width = Inches(3.4)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.05
            r_name = p.add_run(name + '\n')
            r_name.font.name = 'Times New Roman'
            r_name.font.size = Pt(10.0)
            r_name.font.bold = True
            r_affil = p.add_run(affil)
            r_affil.font.name = 'Times New Roman'
            r_affil.font.size = Pt(8.5)
            r_affil.font.italic = False

    # Abstract
    m_abs = re.search(r'\\begin\{abstract\}(.*?)\\end\{abstract\}', tex, re.DOTALL)
    abs_raw = m_abs.group(1).strip() if m_abs else ""
    abs_clean = clean_latex_text(abs_raw, bib_map)

    p_abs = doc.add_paragraph()
    p_abs.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_abs.paragraph_format.space_before = Pt(14)
    p_abs.paragraph_format.space_after = Pt(4)
    p_abs.paragraph_format.line_spacing = 1.05

    r_ab_label = p_abs.add_run('Abstract—')
    r_ab_label.font.name = 'Times New Roman'
    r_ab_label.font.size = Pt(9.0)
    r_ab_label.font.bold = True
    r_ab_label.font.italic = True

    r_ab_text = p_abs.add_run(abs_clean)
    r_ab_text.font.name = 'Times New Roman'
    r_ab_text.font.size = Pt(9.0)
    r_ab_text.font.bold = True

    # Keywords
    m_kw = re.search(r'\\begin\{IEEEkeywords\}(.*?)\\end\{IEEEkeywords\}', tex, re.DOTALL)
    kw_raw = m_kw.group(1).strip() if m_kw else ""
    kw_clean = clean_latex_text(kw_raw, bib_map)

    p_kw = doc.add_paragraph()
    p_kw.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_kw.paragraph_format.space_before = Pt(2)
    p_kw.paragraph_format.space_after = Pt(12)
    p_kw.paragraph_format.line_spacing = 1.05

    r_kw_label = p_kw.add_run('Index Terms—')
    r_kw_label.font.name = 'Times New Roman'
    r_kw_label.font.size = Pt(9.0)
    r_kw_label.font.bold = True
    r_kw_label.font.italic = True

    r_kw_text = p_kw.add_run(kw_clean)
    r_kw_text.font.name = 'Times New Roman'
    r_kw_text.font.size = Pt(9.0)
    r_kw_text.font.bold = True

    # Section 1: 2-Column Section for Body
    s1 = doc.add_section(docx.enum.section.WD_SECTION.CONTINUOUS)
    s1.top_margin = Inches(0.75)
    s1.bottom_margin = Inches(0.75)
    s1.left_margin = Inches(0.75)
    s1.right_margin = Inches(0.75)
    s1.page_width = Inches(8.5)
    s1.page_height = Inches(11.0)

    sectPr = s1._sectPr
    cols = parse_xml(f'<w:cols {nsdecls("w")} w:num="2" w:space="720"/>')
    sectPr.append(cols)

    def add_sec_heading(title):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(title)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(10.0)
        r.font.bold = True

    def add_sub_heading(title):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(title)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(9.5)
        r.font.italic = True
        r.font.bold = True

    def add_body_p(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.05
        r = p.add_run(text)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(9.5)
        return p

    def add_bullet(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.left_indent = Inches(0.2)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.05
        r_b = p.add_run('•  ')
        r_b.font.name = 'Times New Roman'
        r_b.font.size = Pt(9.0)
        r_b.font.bold = True
        r = p.add_run(text)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(9.0)
        return p

    def add_numbered(num, text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.left_indent = Inches(0.2)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.05
        r_n = p.add_run(f'{num}.  ')
        r_n.font.name = 'Times New Roman'
        r_n.font.size = Pt(9.0)
        r_n.font.bold = True
        r = p.add_run(text)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(9.0)
        return p

    def add_figure(img_rel, caption):
        base_dir = os.path.dirname(tex_path)
        img_path = os.path.join(base_dir, img_rel)
        if os.path.exists(img_path):
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.paragraph_format.space_before = Pt(6)
            p_img.paragraph_format.space_after = Pt(2)
            p_img.paragraph_format.keep_with_next = True
            run = p_img.add_run()
            run.add_picture(img_path, width=Inches(3.35))
            
            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap.paragraph_format.space_before = Pt(1)
            p_cap.paragraph_format.space_after = Pt(4)
            p_cap.paragraph_format.keep_with_next = True
            r_cap = p_cap.add_run(caption)
            r_cap.font.name = 'Times New Roman'
            r_cap.font.size = Pt(8.5)
            r_cap.font.italic = True
        else:
            print(f"Warning: image {img_path} not found")

    # Table rendering helpers
    def add_table_1():
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(6)
        p_cap.paragraph_format.space_after = Pt(2)
        p_cap.paragraph_format.keep_with_next = True
        r = p_cap.add_run("TABLE I\nMODEL PERFORMANCE AND CALIBRATION BENCHMARK (N = 2,500, DS-SYNTH-01)")
        r.font.name = 'Times New Roman'
        r.font.size = Pt(8.5)
        r.font.bold = True

        data = [
            ["Model", "Accuracy (S42)", "Mean Acc (5-Seed)", "ROC-AUC", "F1", "ECE (S42)", "Mean ECE", "Brier"],
            ["Logistic Reg.", "99.20%", "98.80% ± 0.40%", "0.9998", "0.9880", "0.0150", "0.0160 ± 0.0030", "0.0100"],
            ["Random Forest", "92.00%", "92.40% ± 0.80%", "0.9825", "0.8878", "0.0710", "0.0720 ± 0.0060", "0.0630"],
            ["Decision Tree", "88.60%", "89.10% ± 1.10%", "0.8710", "0.8350", "0.1140", "0.1120 ± 0.0090", "0.1080"],
            ["MLP (2-layer)", "91.80%", "91.20% ± 1.40%", "0.9750", "0.8820", "0.0820", "0.0850 ± 0.0070", "0.0710"],
            ["XGB (Raw)", "94.60%", "95.20% ± 1.17%", "0.9922", "0.9245", "0.0370", "0.0570 ± 0.0082", "0.0410"],
            ["PRIE (Platt-XGB)", "94.60%", "95.20% ± 1.17%", "0.9922", "0.9245", "0.0212", "0.0350 ± 0.0057", "0.0339"]
        ]
        t = doc.add_table(rows=len(data), cols=len(data[0]))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders_ieee(t)
        for r_i, row in enumerate(data):
            for c_i, val in enumerate(row):
                cell = t.cell(r_i, c_i)
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_i > 0 else WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.space_before = Pt(1)
                p.paragraph_format.space_after = Pt(1)
                run = p.add_run(val)
                run.font.name = 'Times New Roman'
                run.font.size = Pt(7.0)
                if r_i == 0:
                    run.font.bold = True
                    set_header_bottom_border(cell)
                elif r_i == len(data) - 1:
                    run.font.bold = True

    def add_table_2():
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(6)
        p_cap.paragraph_format.space_after = Pt(2)
        p_cap.paragraph_format.keep_with_next = True
        r = p_cap.add_run("TABLE II\nCOUNTERFACTUAL RECOURSE OPTIMIZATION BENCHMARK (EXP-02)")
        r.font.name = 'Times New Roman'
        r.font.size = Pt(8.5)
        r.font.bold = True

        data = [
            ["Recourse Method", "Mean L1 Proximity", "Sparsity (k modified)", "F17 Invariance", "Success Rate"],
            ["Unconstrained DiCE", "0.142 ± 0.028", "5.82 ± 1.12", "12.4%", "98.2%"],
            ["Standard DiCE (No Lock)", "0.188 ± 0.032", "4.10 ± 0.85", "46.8%", "95.5%"],
            ["PRIE Constrained DiCE", "0.214 ± 0.038", "2.47 ± 0.52", "100.0%", "92.8%"]
        ]
        t = doc.add_table(rows=len(data), cols=len(data[0]))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders_ieee(t)
        for r_i, row in enumerate(data):
            for c_i, val in enumerate(row):
                cell = t.cell(r_i, c_i)
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_i > 0 else WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.space_before = Pt(1)
                p.paragraph_format.space_after = Pt(1)
                run = p.add_run(val)
                run.font.name = 'Times New Roman'
                run.font.size = Pt(7.5)
                if r_i == 0:
                    run.font.bold = True
                    set_header_bottom_border(cell)
                elif r_i == len(data) - 1:
                    run.font.bold = True

    def add_table_3():
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(6)
        p_cap.paragraph_format.space_after = Pt(2)
        p_cap.paragraph_format.keep_with_next = True
        r = p_cap.add_run("TABLE III\nDESCRIPTIVE CROSS-STUDY BENCHMARK WITH PRIOR LITERATURE")
        r.font.name = 'Times New Roman'
        r.font.size = Pt(8.5)
        r.font.bold = True

        data = [
            ["Study", "Method", "Reported Accuracy", "Calibration Metric"],
            ["Olipas (2024) [1]", "Random Forest", "88.40%", "Not Evaluated"],
            ["Rao & Swamy (2022) [8]", "Decision Tree", "78.40%", "Not Evaluated"],
            ["Casuat & Festijo (2021) [7]", "Tree Ensemble", "84.50%", "Not Evaluated"],
            ["Patel & Nair (2024) [4]", "Neural Network", "91.20%", "Not Evaluated"],
            ["PRIE (Proposed)", "Platt-XGBoost", "95.20%", "ECE = 0.0350"]
        ]
        t = doc.add_table(rows=len(data), cols=len(data[0]))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders_ieee(t)
        for r_i, row in enumerate(data):
            for c_i, val in enumerate(row):
                cell = t.cell(r_i, c_i)
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_i > 0 else WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.space_before = Pt(1)
                p.paragraph_format.space_after = Pt(1)
                run = p.add_run(val)
                run.font.name = 'Times New Roman'
                run.font.size = Pt(7.5)
                if r_i == 0:
                    run.font.bold = True
                    set_header_bottom_border(cell)
                elif r_i == len(data) - 1:
                    run.font.bold = True

    def add_table_4():
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(6)
        p_cap.paragraph_format.space_after = Pt(2)
        p_cap.paragraph_format.keep_with_next = True
        r = p_cap.add_run("TABLE IV\nMULTIMODAL MOCK INTERVIEW SCORING STABILITY AND LATENCY BENCHMARK (EXP-03)")
        r.font.name = 'Times New Roman'
        r.font.size = Pt(8.5)
        r.font.bold = True

        data = [
            ["Modality / Fusion", "Mean Score", "Variance (σ²)", "Variance Reduction", "Turn Latency"],
            ["Audio Only (Prosody)", "68.42 ± 5.81", "33.76", "Reference", "180 ms"],
            ["Video Only (Face)", "71.18 ± 4.92", "24.21", "-28.29%", "340 ms"],
            ["Speech Only (Whisper)", "74.55 ± 6.24", "38.94", "+15.34%", "450 ms"],
            ["Tri-Modal Late Fusion", "72.48 ± 2.73", "7.43", "-77.98% ± 3.99%", "1,120 ms"]
        ]
        t = doc.add_table(rows=len(data), cols=len(data[0]))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders_ieee(t)
        for r_i, row in enumerate(data):
            for c_i, val in enumerate(row):
                cell = t.cell(r_i, c_i)
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_i > 0 else WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.space_before = Pt(1)
                p.paragraph_format.space_after = Pt(1)
                run = p.add_run(val)
                run.font.name = 'Times New Roman'
                run.font.size = Pt(7.5)
                if r_i == 0:
                    run.font.bold = True
                    set_header_bottom_border(cell)
                elif r_i == len(data) - 1:
                    run.font.bold = True

    # Now parse the sections of paper.tex between \maketitle and \begin{thebibliography}
    body_match = re.search(r'\\maketitle(.*?)\\begin\{thebibliography\}', tex, re.DOTALL)
    if not body_match:
        print("Error: Could not isolate body between \\maketitle and \\begin{thebibliography}")
        return

    body_raw = body_match.group(1)

    # We iterate line by line or block by block
    lines = body_raw.split('\n')
    idx = 0
    in_abstract = False
    in_keywords = False
    in_enumerate = False
    in_itemize = False
    enum_counter = 0

    while idx < len(lines):
        line = lines[idx].strip()
        idx += 1

        if not line or line.startswith('%'):
            continue

        # Skip abstract and keywords since they were processed in header
        if r'\begin{abstract}' in line:
            in_abstract = True
            continue
        if r'\end{abstract}' in line:
            in_abstract = False
            continue
        if r'\begin{IEEEkeywords}' in line:
            in_keywords = True
            continue
        if r'\end{IEEEkeywords}' in line:
            in_keywords = False
            continue
        if in_abstract or in_keywords:
            continue

        # Check section headings
        m_sec = re.match(r'\\section\*?\{([^}]+)\}', line)
        if m_sec:
            sec_title = m_sec.group(1).strip()
            # If Roman numeral not present, compute Roman numeral
            sec_upper = sec_title.upper()
            if sec_title == "Acknowledgment":
                add_sec_heading("ACKNOWLEDGMENT")
            elif not re.match(r'^[IVXLCDM]+\.', sec_upper):
                # map known sections
                sec_map = {
                    "INTRODUCTION": "I. INTRODUCTION",
                    "RELATED WORK AND RESEARCH GAP": "II. RELATED WORK AND RESEARCH GAP",
                    "PRIE SYSTEM ARCHITECTURE": "III. PRIE SYSTEM ARCHITECTURE",
                    "METHODOLOGY": "IV. METHODOLOGY",
                    "EXPERIMENTAL DESIGN": "V. EXPERIMENTAL DESIGN",
                    "RESULTS": "VI. RESULTS",
                    "DISCUSSION": "VII. DISCUSSION",
                    "LIMITATIONS AND THREATS TO VALIDITY": "VIII. LIMITATIONS AND THREATS TO VALIDITY",
                    "CONCLUSION": "IX. CONCLUSION"
                }
                add_sec_heading(sec_map.get(sec_upper, sec_upper))
            else:
                add_sec_heading(sec_upper)
            continue

        # Check subsection headings
        m_subsec = re.match(r'\\subsection\{([^}]+)\}', line)
        if m_subsec:
            subsec_title = m_subsec.group(1).strip()
            add_sub_heading(clean_latex_text(subsec_title, bib_map))
            continue

        # Check subsubsection headings
        m_subsubsec = re.match(r'\\subsubsection\{([^}]+)\}', line)
        if m_subsubsec:
            subsubsec_title = m_subsubsec.group(1).strip()
            add_sub_heading(clean_latex_text(subsubsec_title, bib_map))
            continue

        # Check list environments
        if r'\begin{enumerate}' in line:
            in_enumerate = True
            enum_counter = 0
            continue
        if r'\end{enumerate}' in line:
            in_enumerate = False
            continue
        if r'\begin{itemize}' in line:
            in_itemize = True
            continue
        if r'\end{itemize}' in line:
            in_itemize = False
            continue

        # Check list items
        if line.startswith(r'\item'):
            item_text = line[5:].strip()
            # Accumulate multi-line item
            while idx < len(lines) and not lines[idx].strip().startswith(r'\item') and not lines[idx].strip().startswith(r'\end{') and not lines[idx].strip().startswith(r'\section') and not lines[idx].strip().startswith(r'\subsection'):
                nxt = lines[idx].strip()
                idx += 1
                if nxt and not nxt.startswith('%'):
                    item_text += " " + nxt

            cleaned_item = clean_latex_text(item_text, bib_map)
            if in_enumerate:
                enum_counter += 1
                add_numbered(enum_counter, cleaned_item)
            else:
                add_bullet(cleaned_item)
            continue

        # Check figures
        if r'\begin{figure}' in line:
            # gather figure lines
            fig_lines = []
            while idx < len(lines) and r'\end{figure}' not in lines[idx]:
                fig_lines.append(lines[idx])
                idx += 1
            idx += 1 # skip \end{figure}
            fig_text = "\n".join(fig_lines)
            m_img = re.search(r'\\includegraphics\[.*?\]\{([^}]+)\}', fig_text)
            m_cap = re.search(r'\\caption\{([^}]+)\}', fig_text)
            if m_img and m_cap:
                img_rel = m_img.group(1).strip()
                cap_clean = clean_latex_text(m_cap.group(1).strip(), bib_map)
                add_figure(img_rel, cap_clean)
            continue

        # Check tables
        if r'\begin{table}' in line:
            tbl_lines = []
            while idx < len(lines) and r'\end{table}' not in lines[idx]:
                tbl_lines.append(lines[idx])
                idx += 1
            idx += 1 # skip \end{table}
            tbl_text = "\n".join(tbl_lines)
            if 'tab1' in tbl_text or 'MODEL PERFORMANCE' in tbl_text:
                add_table_1()
            elif 'tab2' in tbl_text or 'COUNTERFACTUAL' in tbl_text:
                add_table_2()
            elif 'tab3' in tbl_text or 'CROSS-STUDY' in tbl_text:
                add_table_3()
            elif 'tab4' in tbl_text or 'MULTIMODAL' in tbl_text:
                add_table_4()
            continue

        # Skip display equations (render them cleanly if needed, or represent them as math equations)
        if r'\begin{equation}' in line:
            eq_lines = []
            while idx < len(lines) and r'\end{equation}' not in lines[idx]:
                eq_lines.append(lines[idx])
                idx += 1
            idx += 1 # skip \end{equation}
            eq_text = " ".join([l.strip() for l in eq_lines if l.strip()])
            cleaned_eq = clean_latex_text(eq_text, bib_map)
            p_eq = doc.add_paragraph()
            p_eq.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_eq.paragraph_format.space_before = Pt(3)
            p_eq.paragraph_format.space_after = Pt(3)
            r_eq = p_eq.add_run(cleaned_eq)
            r_eq.font.name = 'Times New Roman'
            r_eq.font.size = Pt(9.0)
            r_eq.font.italic = True
            continue

        # Normal paragraph text: accumulate till blank line
        para_lines = [line]
        while idx < len(lines):
            nxt = lines[idx].strip()
            if not nxt:
                idx += 1
                break
            if nxt.startswith(r'\section') or nxt.startswith(r'\subsection') or nxt.startswith(r'\subsubsection') or nxt.startswith(r'\begin{') or nxt.startswith(r'\item'):
                break
            if not nxt.startswith('%'):
                para_lines.append(nxt)
            idx += 1

        para_raw = " ".join(para_lines)
        cleaned_p = clean_latex_text(para_raw, bib_map)
        if cleaned_p:
            add_body_p(cleaned_p)

    # References Section
    add_sec_heading("REFERENCES")
    for idx_b, key_b, body_b in bib_items:
        p_ref = doc.add_paragraph()
        p_ref.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p_ref.paragraph_format.left_indent = Inches(0.2)
        p_ref.paragraph_format.first_line_indent = Inches(-0.2)
        p_ref.paragraph_format.space_before = Pt(0)
        p_ref.paragraph_format.space_after = Pt(2)
        p_ref.paragraph_format.line_spacing = 1.05

        r_num = p_ref.add_run(f'[{idx_b}]  ')
        r_num.font.name = 'Times New Roman'
        r_num.font.size = Pt(8.0)

        r_txt = p_ref.add_run(body_b)
        r_txt.font.name = 'Times New Roman'
        r_txt.font.size = Pt(8.0)

    # Save document
    print(f"Saving {docx_out}...")
    doc.save(docx_out)

    # Extract text to txt_out
    print(f"Exporting text to {txt_out}...")
    with open(txt_out, 'w', encoding='utf-8') as f:
        f.write("=====================================================================\n")
        f.write("PRIE: Placement Readiness Intelligence Engine (IEEE Manuscript Text)\n")
        f.write("=====================================================================\n\n")
        for p in doc.paragraphs:
            txt = p.text.strip()
            if txt:
                f.write(txt + "\n\n")
        for t in doc.tables:
            f.write("--- TABLE DATA ---\n")
            for r in t.rows:
                row_vals = [c.text.strip().replace('\n', ' ') for c in r.cells]
                f.write(" | ".join(row_vals) + "\n")
            f.write("\n")

    print("Synchronization completed successfully!")

if __name__ == '__main__':
    base = r"d:\4-1 AD\All College Docs and ppts\Documentations\PDR\PRIE-Research\10_Publication\Conference_Paper"
    tex_file = os.path.join(base, "paper.tex")
    docx_file = os.path.join(base, "paper.docx")
    txt_file = os.path.join(base, "paper.docx.txt")
    convert_tex_to_docx(tex_file, docx_file, txt_file)

    # Also mirror to 01_Conference_Paper
    base_01 = r"d:\4-1 AD\All College Docs and ppts\Documentations\PDR\PRIE-Research\10_Publication\01_Conference_Paper"
    shutil.copyfile(docx_file, os.path.join(base_01, "paper.docx"))
    shutil.copyfile(txt_file, os.path.join(base_01, "paper.docx.txt"))
    print("Mirrored to 01_Conference_Paper!")
